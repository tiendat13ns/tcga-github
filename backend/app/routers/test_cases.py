from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy import or_, update
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
import io

from app.core.auth import get_current_user, get_db
from app.core.ownership import verify_requirement_owner
from app.models import Requirement, User
from app.schemas.requirement_schema import GenerationStartedResponse
from app.schemas.test_case_schema import ListTestCasesResponse
from app.services.credit_service import CREDIT_COST
from app.services.generation import job_runner
from app.services.generation.test_case_generation_service import (
    TestCaseGenerationError,
    list_test_cases_by_requirement,
    run_test_case_generation_job,
)

router = APIRouter(prefix="/api/v1/requirements", tags=["test-cases"])


@router.post(
    "/{requirement_id}/test-cases/generate",
    response_model=GenerationStartedResponse,
    status_code=202,
)
async def generate_test_cases(
    requirement_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Nhận yêu cầu sinh test case rồi chạy NỀN — trả 202 ngay. Frontend poll trạng thái qua
    requirement (test_case_status). Credit chỉ trừ khi job thành công (trong background)."""
    req = verify_requirement_owner(db, requirement_id, current_user)
    # Chặn sớm nếu không đủ credit — tránh chạy job rồi mới báo thiếu.
    cost = CREDIT_COST.get("TEST_CASE_GENERATION", 0)
    if current_user.credit_balance < cost:
        raise HTTPException(
            status_code=402,
            detail=f"Không đủ Credit. Cần {cost} Credits nhưng chỉ còn {current_user.credit_balance}.",
        )
    # UPDATE có điều kiện (thay vì check-rồi-set riêng lẻ) để tránh race: hai request đến gần
    # như đồng thời chỉ có 1 cái khớp WHERE và thắng, cái còn lại rowcount=0 -> 409. Nhờ vậy
    # không thể có 2 job nền chạy song song cho cùng 1 requirement (double LLM call, double credit).
    result = db.execute(
        update(Requirement)
        .where(Requirement.id == req.id)
        .where(or_(Requirement.test_case_status.is_(None), Requirement.test_case_status != "generating"))
        .values(test_case_status="generating", test_case_error=None)
    )
    db.commit()
    if result.rowcount == 0:
        raise HTTPException(status_code=409, detail="Requirement này đang được sinh test case.")

    job_runner.submit(run_test_case_generation_job(requirement_id, str(current_user.id)))
    return GenerationStartedResponse(status="generating", requirement_id=requirement_id)


@router.get(
    "/{requirement_id}/test-cases",
    response_model=ListTestCasesResponse,
)
def get_test_cases(
    requirement_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    verify_requirement_owner(db, requirement_id, current_user)
    try:
        return list_test_cases_by_requirement(requirement_id)
    except TestCaseGenerationError as exc:
        raise HTTPException(status_code=exc.status_code, detail=str(exc)) from exc
    except SQLAlchemyError as exc:
        raise HTTPException(status_code=500, detail="Database error while loading test cases") from exc


@router.get(
    "/{requirement_id}/test-cases/export",
    responses={
        200: {
            "content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}},
            "description": "Exported Excel file containing test cases.",
        }
    },
)
def export_test_cases(
    requirement_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    verify_requirement_owner(db, requirement_id, current_user)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        
        # Load test cases from DB
        response_data = list_test_cases_by_requirement(requirement_id)
        tcs = response_data.test_cases
        
        # We need the requirement to get the Feature name (functional_requirement)
        # For simplicity, we can fetch it here or just leave it empty if we don't have the repo
        # Let's import SessionLocal and Requirement to get the feature name
        from app.database import SessionLocal
        from app.models import Requirement
        from uuid import UUID
        
        feature_name = "Unknown Feature"
        try:
            with SessionLocal() as db:
                req = db.get(Requirement, UUID(requirement_id))
                if req:
                    feature_name = req.feature_name or req.module_name or req.title or "Unknown Feature"
        except Exception:
            pass

        wb = Workbook()
        ws = wb.active
        ws.title = "Test Cases"
        headers = ["Feature", "Test Case ID", "Test Item", "Precondition", "Test Steps", "Test Data", "Expected Output"]
        ws.append(headers)
        
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            
        for idx, tc in enumerate(tcs):
            # Format test steps as a numbered string
            steps_text = "\n".join([f"{i+1}. {s}" for i, s in enumerate(tc.test_steps)]) if tc.test_steps else ""
            
            ws.append([
                feature_name,
                f"TC-{str(idx+1).zfill(2)}",
                tc.title,
                tc.preconditions or "",
                steps_text,
                tc.test_data or "",
                tc.expected_result
            ])
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="test_cases_{requirement_id[:8]}.xlsx"'
            }
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc
