import io
from uuid import UUID
from fastapi import APIRouter, Depends, HTTPException, Response, Query
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session
from datetime import datetime, timezone
import logging

from app.core.auth import get_current_user, get_db
from app.core.ownership import verify_requirement_owner
from app.models import Document, Project, TestCase, Requirement, User
from app.schemas.test_case_schema import (
    StudioTestCaseItem,
    StudioTestCaseListResponse,
    TestCaseUpdatePayload,
    TestCaseCreatePayload,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/test-cases", tags=["test-case-studio"])


def _get_owned_project_id(db: Session, project_id: str, user: User) -> UUID:
    """Verify project_id hợp lệ và thuộc về user hiện tại — trả 404 để tránh lộ sự tồn tại
    của project cho user khác."""
    try:
        project_uuid = UUID(project_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid project_id format") from exc

    project = db.get(Project, project_uuid)
    if project is None or project.user_id != user.id:
        raise HTTPException(status_code=404, detail="Project not found")
    return project_uuid


def _verify_test_case_owner(db: Session, test_case_id: str, user: User) -> TestCase:
    """Đảm bảo test case thuộc project của chính user đang đăng nhập (qua requirement -> project)."""
    try:
        tc_uuid = UUID(test_case_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Invalid test_case_id format") from exc

    tc = db.query(TestCase).filter(TestCase.id == tc_uuid).first()
    if tc is None:
        raise HTTPException(status_code=404, detail="Test case not found")

    try:
        verify_requirement_owner(db, str(tc.requirement_id), user)
    except HTTPException:
        raise HTTPException(status_code=404, detail="Test case not found")
    return tc


def _to_studio_item(tc: TestCase, req: Requirement) -> StudioTestCaseItem:
    return StudioTestCaseItem(
        id=str(tc.id),
        requirement_id=str(tc.requirement_id),
        document_id=str(tc.document_id) if tc.document_id else None,
        title=tc.title,
        scenario=tc.scenario,
        preconditions=tc.preconditions,
        test_steps=tc.test_steps,
        test_data=tc.test_data,
        expected_result=tc.expected_result,
        actual_result=tc.actual_result,
        priority=tc.priority,
        severity=tc.severity,
        test_type=tc.test_type,
        automation_candidate=tc.automation_candidate,
        execution_type=tc.execution_type,
        execution_status=tc.execution_status,
        status=tc.status,
        note=tc.note,
        version=tc.version,
        feature_name=req.feature_name,
        requirement_title=req.title,
        project_id=str(req.project_id) if req.project_id else None,
        module_name=req.module_name,
    )

@router.get("/export",
    responses={
        200: {
            "content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}},
            "description": "Exported Excel file containing test cases.",
        }
    },
)
def export_test_cases_studio(
    project_id: str | None = None,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if not project_id:
        raise HTTPException(status_code=400, detail="project_id is required")
    project_uuid = _get_owned_project_id(db, project_id, current_user)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        query = db.query(TestCase, Requirement).join(Requirement, TestCase.requirement_id == Requirement.id)
        query = query.filter(Requirement.project_id == project_uuid)
        query = query.order_by(TestCase.created_at.asc(), TestCase.id.asc())
        results = query.all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Test Cases"
        headers = ["Feature", "Test Case ID", "Title", "Precondition", "Test Steps", "Test Data", "Expected Output", "Priority", "Note", "Test Type"]
        ws.append(headers)

        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill

        for idx, (tc, req) in enumerate(results):
            feature_name = req.feature_name or req.module_name or req.title or "Unknown Feature"
            steps_text = "\n".join([f"{i+1}. {s}" for i, s in enumerate(tc.test_steps)]) if tc.test_steps else ""

            ws.append([
                feature_name,
                f"TC-{str(idx+1).zfill(2)}",
                tc.title,
                tc.preconditions or "",
                steps_text,
                tc.test_data or "",
                tc.expected_result,
                tc.priority,
                tc.note or "",
                tc.test_type or ""
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"test_cases_{project_id[:8]}.xlsx"

        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"'
            }
        )
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.get("", response_model=StudioTestCaseListResponse)
def get_studio_test_cases(
    project_id: str | None = Query(None),
    document_id: str | None = Query(None),
    priority: str | None = Query(None),
    status: str | None = Query(None),
    test_type: str | None = Query(None),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    try:
        query = db.query(TestCase, Requirement).join(Requirement, TestCase.requirement_id == Requirement.id)

        if project_id:
            project_uuid = _get_owned_project_id(db, project_id, current_user)
            query = query.filter(Requirement.project_id == project_uuid)
        else:
            # Không truyền project_id: chỉ trả về test case thuộc các project của user hiện tại.
            query = query.join(Project, Requirement.project_id == Project.id).filter(Project.user_id == current_user.id)

        if document_id:
            try:
                query = query.filter(TestCase.document_id == UUID(document_id))
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid document_id format")
        if priority:
            query = query.filter(TestCase.priority == priority)
        if status:
            query = query.filter(TestCase.status == status)
        if test_type:
            query = query.filter(TestCase.test_type == test_type)

        query = query.order_by(TestCase.created_at.asc(), TestCase.id.asc())
        results = query.all()
        items = [_to_studio_item(tc, req) for tc, req in results]

        return StudioTestCaseListResponse(
            total_test_cases=len(items),
            test_cases=items
        )
    except HTTPException:
        raise
    except SQLAlchemyError as exc:
        logger.error(f"DB Error: {exc}")
        raise HTTPException(status_code=500, detail="Database error while fetching test cases") from exc


@router.put("/{test_case_id}", response_model=StudioTestCaseItem)
def update_studio_test_case(
    test_case_id: str,
    payload: TestCaseUpdatePayload,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tc = _verify_test_case_owner(db, test_case_id, current_user)
    try:
        update_data = payload.model_dump(exclude_unset=True)
        if update_data:
            for key, value in update_data.items():
                setattr(tc, key, value)
            tc.updated_at = datetime.now(timezone.utc)
            tc.version += 1

            db.commit()
            db.refresh(tc)

        req = db.query(Requirement).filter(Requirement.id == tc.requirement_id).first()
        if not req:
            raise HTTPException(status_code=500, detail="Requirement not found for test case")

        return _to_studio_item(tc, req)
    except HTTPException:
        raise
    except SQLAlchemyError as exc:
        logger.error(f"DB Error: {exc}")
        raise HTTPException(status_code=500, detail="Database error while updating test case") from exc

@router.post("", response_model=StudioTestCaseItem)
def create_studio_test_case(
    payload: TestCaseCreatePayload,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    req = verify_requirement_owner(db, payload.requirement_id, current_user)
    req_uuid = req.id

    try:
        tc = TestCase(
            requirement_id=req_uuid,
            document_id=req.document_id,
            title=payload.title,
            scenario=payload.scenario,
            preconditions=payload.preconditions,
            test_steps=payload.test_steps,
            test_data=payload.test_data,
            expected_result=payload.expected_result,
            actual_result=payload.actual_result,
            priority=payload.priority,
            severity=payload.severity,
            test_type=payload.test_type,
            automation_candidate=payload.automation_candidate,
            execution_type=payload.execution_type,
            execution_status=payload.execution_status,
            status=payload.status,
            note=payload.note,
        )
        db.add(tc)
        db.commit()
        db.refresh(tc)

        return _to_studio_item(tc, req)
    except HTTPException:
        raise
    except SQLAlchemyError as exc:
        logger.error(f"DB Error: {exc}")
        raise HTTPException(status_code=500, detail="Database error while creating test case") from exc

from pydantic import BaseModel
class BugReportPayload(BaseModel):
    actual_result: str

@router.post("/{test_case_id}/bug-report")
def generate_auto_bug_report(
    test_case_id: str,
    payload: BugReportPayload,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    tc = _verify_test_case_owner(db, test_case_id, current_user)
    try:
        tc_data = {
            "title": tc.title,
            "preconditions": tc.preconditions,
            "test_steps": tc.test_steps,
            "expected_result": tc.expected_result
        }

        from app.services.agent.bug_report_service import generate_bug_report
        report_content = generate_bug_report(tc_data, payload.actual_result)

        return {"report": report_content}
    except HTTPException:
        raise
    except Exception as exc:
        logger.error(f"Error generating bug report: {exc}")
        raise HTTPException(status_code=500, detail=str(exc)) from exc
